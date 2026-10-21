"""生成 samples/ 里的虚构演示材料。

为什么要有这个脚本：演示数据必须可复现、可审查，而不是手工做三份文件塞进仓库。
每次运行都会重写 samples/ 下的三个文件，内容里的冲突点是故意埋的，
用来验证 audit.py 能抓到「字符转置」「取值不一致」「格式违规」三类问题。

全部信息均为虚构。
"""

from __future__ import annotations

import os

from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
SAMPLES = os.path.join(os.path.dirname(HERE), "samples")

RESUME_LINES = [
    ("林小舟", "title"),
    ("手机号：13800001234", "p"),
    ("电子邮箱：linxz@example.com", "p"),
    ("出生年月：2002.11", "p"),
    ("预计毕业时间：2027.06", "p"),
    ("政治面貌：中共党员", "p"),
    ("", "p"),
    ("教育经历", "h"),
    ("东岭大学　图书情报　硕士　2024.09 – 2027.06", "p"),
    ("南湖师范大学　信息资源管理　本科　2020.09 – 2024.06", "p"),
    ("", "p"),
    ("成果", "h"),
    ("软件著作权登记号：2024SR1234567", "p"),
    ("软件著作权名称：脑电信号分析系统 V1.0", "p"),
    ("", "p"),
    ("材料填报口径", "h"),
    ("本科毕业院校：南湖师范大学", "p"),
    ("本科所学专业：信息资源管理", "p"),
    ("硕士毕业院校：东岭大学", "p"),
    ("硕士所学专业：图书情报", "p"),
]

MASTER_MD = """# 网申字段母版（示例）

姓名：林小舟
手机号：13800001234
电子邮箱：linxz@example.com
出生年月：2002.11
毕业时间：2027年7月
本科院校：南湖师范大学
本科专业：信息管理与信息系统
硕士院校：东岭大学
硕士专业：图书情报
软著登记号：2024SR1234567
"""

CHECKLIST_MD = """# 证明材料清单（示例）

1. 学历学位_本科_教育部学历证书电子注册备案表.pdf
2. 学历学位_硕士_教育部学籍在线验证报告.pdf
3. 软著登记证书.pdf —— 登记号 2024SR1234576
4. 英语四级成绩报告单.pdf
5. 回执联系手机号：1380001234
"""


def build() -> list[str]:
    os.makedirs(SAMPLES, exist_ok=True)
    made: list[str] = []

    doc = Document()
    for text, kind in RESUME_LINES:
        if not text:
            doc.add_paragraph("")
        elif kind == "title":
            doc.add_heading(text, level=0)
        elif kind == "h":
            doc.add_heading(text, level=2)
        else:
            doc.add_paragraph(text)
    p = os.path.join(SAMPLES, "01_简历_林小舟.docx")
    doc.save(p)
    made.append(p)

    for name, body in [
        ("02_网申字段母版.md", MASTER_MD),
        ("03_证明材料清单.md", CHECKLIST_MD),
    ]:
        q = os.path.join(SAMPLES, name)
        with open(q, "w", encoding="utf-8") as f:
            f.write(body)
        made.append(q)

    return made


if __name__ == "__main__":
    for f in build():
        print("已生成", os.path.relpath(f, os.path.dirname(HERE)), os.path.getsize(f), "B")
